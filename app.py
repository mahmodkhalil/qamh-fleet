from flask import Flask, render_template, request, redirect, url_for, flash, send_file, session, send_from_directory, jsonify
from werkzeug.utils import secure_filename
from functools import wraps
import sqlite3
import json
import sys
import subprocess
import os
import base64
from datetime import datetime, timedelta
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
import requests

try:
    import psycopg2
    from psycopg2.extras import RealDictCursor
    HAS_POSTGRES = True
except ImportError:
    HAS_POSTGRES = False

app = Flask(__name__)
app.secret_key = 'qamh_fleet_secret_2026'

DATA_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(DATA_DIR, 'qamh_fleet.db')
UPLOAD_FOLDER = os.path.join(DATA_DIR, 'uploads')
PASSWORDS_FILE = os.path.join(DATA_DIR, 'passwords.json')
ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif', 'pdf'}
DATABASE_URL = os.environ.get('DATABASE_URL')
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['MAX_CONTENT_LENGTH'] = 32 * 1024 * 1024

if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)

TELEGRAM_TOKEN = os.environ.get('TELEGRAM_TOKEN', 'ضع_توكن_البوت_هنا')
GEMINI_API_KEY = os.environ.get('GEMINI_API_KEY', 'ضع_مفتاح_جيميني_هنا')
SUPABASE_URL = os.environ.get('SUPABASE_URL', '')
SUPABASE_SECRET_KEY = os.environ.get('SUPABASE_SECRET_KEY', '')
SUPABASE_BUCKET = 'qamh-images'

def get_connection():
    if DATABASE_URL: return psycopg2.connect(DATABASE_URL, cursor_factory=RealDictCursor)
    else:
        conn = sqlite3.connect(DB_PATH); conn.row_factory = sqlite3.Row; return conn

def is_postgres(): return DATABASE_URL is not None
def execute_query(cursor, query, params=()):
    if is_postgres(): query = query.replace('?', '%s')
    cursor.execute(query, params); return cursor

def is_admin(): return session.get('role') == 'admin'
def is_viewer(): return session.get('role') in ('admin', 'viewer')

def load_passwords():
    if not os.path.exists(PASSWORDS_FILE):
        default = {'admin_password': 'admin2026', 'viewer_password': 'qamh2026'}
        with open(PASSWORDS_FILE, 'w', encoding='utf-8') as f: json.dump(default, f, ensure_ascii=False, indent=4)
        return default
    try:
        with open(PASSWORDS_FILE, 'r', encoding='utf-8') as f: return json.load(f)
    except Exception: return {'admin_password': 'admin2026', 'viewer_password': 'qamh2026'}

def save_passwords(data):
    with open(PASSWORDS_FILE, 'w', encoding='utf-8') as f: json.dump(data, f, ensure_ascii=False, indent=4)

def get_admin_password(): return load_passwords().get('admin_password', 'admin2026')
def get_viewer_password(): return load_passwords().get('viewer_password', 'qamh2026')
def allowed_file(filename): return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not is_viewer(): return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function

def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not is_admin(): flash('هذا الإجراء متاح للمسؤول فقط', 'error'); return redirect(url_for('index'))
        return f(*args, **kwargs)
    return decorated_function

def get_effective_cost(f):
    if f['source'] == 'خزان الشركة': return 0
    ps = f['payment_status'] or ''
    if ps not in ('مدفوعه من قسم الحركه', 'مدفوعة من قسم الحركة'): return 0
    return f['final_cost'] if f['final_cost'] else (f['total_cost'] or 0)

def get_maintenance_alerts():
    conn = get_connection(); cursor = conn.cursor()
    execute_query(cursor, "SELECT m.*, c.name as car_name FROM maintenance m JOIN cars c ON m.car_id = c.id WHERE m.next_service_km IS NOT NULL AND m.category = 'دورية' ORDER BY m.car_id, m.km_at_service DESC")
    all_maint = cursor.fetchall(); conn.close(); latest = {}
    for m in all_maint:
        if m['car_id'] not in latest: latest[m['car_id']] = m
    alerts = []
    for car_id, m in latest.items():
        conn2 = get_connection(); cur2 = conn2.cursor()
        execute_query(cur2, "SELECT MAX(end_km) as max_km FROM trips WHERE car_id=? AND end_km IS NOT NULL", (car_id,))
        row = cur2.fetchone(); trip_km = row['max_km'] if row else None
        execute_query(cur2, "SELECT MAX(km_at_service) as max_km FROM maintenance WHERE car_id=? AND km_at_service IS NOT NULL", (car_id,))
        row = cur2.fetchone(); maint_km = row['max_km'] if row else None
        conn2.close()
        current_km = max([x for x in [trip_km, maint_km] if x is not None]) if (trip_km or maint_km) else 0
        remaining = m['next_service_km'] - current_km
        if remaining <= 0: status = 'red'; message = f"تجاوزت الصيانة بـ {abs(remaining):.0f} كم"
        elif remaining <= 500: status = 'yellow'; message = f"باقي {remaining:.0f} كم"
        else: continue
        alerts.append({'car_name': m['car_name'], 'status': status, 'message': message, 'current_km': current_km, 'next_km': m['next_service_km'], 'remaining': remaining})
    return alerts

def ensure_all_tables():
    conn = get_connection(); cursor = conn.cursor()
    if is_postgres():
        execute_query(cursor, "CREATE TABLE IF NOT EXISTS cars (id SERIAL PRIMARY KEY, name TEXT, plate_number TEXT, car_type TEXT, notes TEXT, active INTEGER DEFAULT 1, expected_consumption REAL DEFAULT 0, fuel_type TEXT DEFAULT '')")
        execute_query(cursor, "CREATE TABLE IF NOT EXISTS drivers (id SERIAL PRIMARY KEY, name TEXT, phone TEXT, notes TEXT, active INTEGER DEFAULT 1)")
        execute_query(cursor, "CREATE TABLE IF NOT EXISTS requesting_parties (id SERIAL PRIMARY KEY, name TEXT, notes TEXT, active INTEGER DEFAULT 1)")
        execute_query(cursor, "CREATE TABLE IF NOT EXISTS trips (id SERIAL PRIMARY KEY, trip_date TEXT, car_id INTEGER, driver_id INTEGER, destination TEXT, start_km REAL, end_km REAL, distance REAL, notes TEXT, trip_source TEXT DEFAULT 'manual', admin_seen INTEGER DEFAULT 0)")
        execute_query(cursor, "CREATE TABLE IF NOT EXISTS fuel (id SERIAL PRIMARY KEY, fuel_date TEXT, car_id INTEGER, driver_id INTEGER, fuel_type TEXT, source TEXT, payment_status TEXT, liters REAL, price_per_liter REAL, total_cost REAL, final_cost REAL, km_at_fill REAL, notes TEXT)")
        execute_query(cursor, "CREATE TABLE IF NOT EXISTS maintenance (id SERIAL PRIMARY KEY, date TEXT, car_id INTEGER, type TEXT, category TEXT, km_at_service REAL, cost REAL, invoice_number TEXT, invoice_image TEXT, workshop TEXT, notes TEXT, next_service_km REAL, next_service_date TEXT, currency TEXT DEFAULT 'ل.س')")
        execute_query(cursor, "CREATE TABLE IF NOT EXISTS daily_checks (id SERIAL PRIMARY KEY, check_date TEXT, check_time TEXT, car_id INTEGER, driver_id INTEGER, km_at_check REAL, general_notes TEXT, has_issue INTEGER DEFAULT 0, created_at TEXT)")
        execute_query(cursor, "CREATE TABLE IF NOT EXISTS pending_fuel (id SERIAL PRIMARY KEY, driver_id INTEGER, car_id INTEGER, fuel_date TEXT, fuel_time TEXT, liters REAL, km_at_fill REAL, station_image TEXT, car_image TEXT, notes TEXT, status TEXT DEFAULT 'pending', created_at TEXT)")
        execute_query(cursor, "CREATE TABLE IF NOT EXISTS travel_missions (id SERIAL PRIMARY KEY, date TEXT, driver_id INTEGER, description TEXT, distance_km REAL DEFAULT 0, is_travel INTEGER DEFAULT 0, notes TEXT)")
        execute_query(cursor, "CREATE TABLE IF NOT EXISTS overtime (id SERIAL PRIMARY KEY, date TEXT, driver_id INTEGER, morning_hours REAL DEFAULT 0, evening_hours REAL DEFAULT 0, notes TEXT)")
    else:
        cursor.execute("CREATE TABLE IF NOT EXISTS cars (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT, plate_number TEXT, car_type TEXT, notes TEXT, active INTEGER DEFAULT 1, expected_consumption REAL DEFAULT 0, fuel_type TEXT DEFAULT '')")
        cursor.execute("CREATE TABLE IF NOT EXISTS drivers (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT, phone TEXT, notes TEXT, active INTEGER DEFAULT 1)")
        cursor.execute("CREATE TABLE IF NOT EXISTS requesting_parties (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT, notes TEXT, active INTEGER DEFAULT 1)")
        cursor.execute("CREATE TABLE IF NOT EXISTS trips (id INTEGER PRIMARY KEY AUTOINCREMENT, trip_date TEXT, car_id INTEGER, driver_id INTEGER, destination TEXT, start_km REAL, end_km REAL, distance REAL, notes TEXT, trip_source TEXT DEFAULT 'manual', admin_seen INTEGER DEFAULT 0)")
        cursor.execute("CREATE TABLE IF NOT EXISTS fuel (id INTEGER PRIMARY KEY AUTOINCREMENT, fuel_date TEXT, car_id INTEGER, driver_id INTEGER, fuel_type TEXT, source TEXT, payment_status TEXT, liters REAL, price_per_liter REAL, total_cost REAL, final_cost REAL, km_at_fill REAL, notes TEXT)")
        cursor.execute("CREATE TABLE IF NOT EXISTS maintenance (id INTEGER PRIMARY KEY AUTOINCREMENT, date TEXT, car_id INTEGER, type TEXT, category TEXT, km_at_service REAL, cost REAL, invoice_number TEXT, invoice_image TEXT, workshop TEXT, notes TEXT, next_service_km REAL, next_service_date TEXT, currency TEXT DEFAULT 'ل.س')")
        cursor.execute("CREATE TABLE IF NOT EXISTS daily_checks (id INTEGER PRIMARY KEY AUTOINCREMENT, check_date TEXT, check_time TEXT, car_id INTEGER, driver_id INTEGER, km_at_check REAL, general_notes TEXT, has_issue INTEGER DEFAULT 0, created_at TEXT)")
        cursor.execute("CREATE TABLE IF NOT EXISTS pending_fuel (id INTEGER PRIMARY KEY AUTOINCREMENT, driver_id INTEGER, car_id INTEGER, fuel_date TEXT, fuel_time TEXT, liters REAL, km_at_fill REAL, station_image TEXT, car_image TEXT, notes TEXT, status TEXT DEFAULT 'pending', created_at TEXT)")
        cursor.execute("CREATE TABLE IF NOT EXISTS travel_missions (id INTEGER PRIMARY KEY AUTOINCREMENT, date TEXT, driver_id INTEGER, description TEXT, distance_km REAL DEFAULT 0, is_travel INTEGER DEFAULT 0, notes TEXT)")
        cursor.execute("CREATE TABLE IF NOT EXISTS overtime (id INTEGER PRIMARY KEY AUTOINCREMENT, date TEXT, driver_id INTEGER, morning_hours REAL DEFAULT 0, evening_hours REAL DEFAULT 0, notes TEXT)")
    conn.commit(); conn.close()

ensure_all_tables()

# ============== المسارات ==============
@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        password = request.form.get('password', ''); login_type = request.form.get('login_type', 'viewer')
        if login_type == 'admin':
            if password == get_admin_password(): session['role'] = 'admin'; return redirect(url_for('index'))
            else: flash('خطأ', 'error')
        else:
            if password == get_viewer_password(): session['role'] = 'viewer'; return redirect(url_for('index'))
            else: flash('خطأ', 'error')
    return render_template('login.html')

@app.route('/logout')
def logout(): session.clear(); return redirect(url_for('login'))

@app.route('/')
@login_required
def index():
    conn = get_connection(); cursor = conn.cursor()
    execute_query(cursor, "SELECT COUNT(*) as cnt FROM cars WHERE active=1"); cars_count = cursor.fetchone()['cnt']
    execute_query(cursor, "SELECT COUNT(*) as cnt FROM drivers WHERE active=1"); drivers_count = cursor.fetchone()['cnt']
    execute_query(cursor, "SELECT t.*, c.name as car_name, d.name as driver_name FROM trips t JOIN cars c ON t.car_id = c.id JOIN drivers d ON t.driver_id = d.id ORDER BY t.id DESC LIMIT 10")
    recent = cursor.fetchall(); conn.close()
    stats = {'cars': cars_count, 'drivers': drivers_count, 'trips_month': 0, 'km_month': 0}
    return render_template('index.html', stats=stats, trips_by_car=recent, alerts=get_maintenance_alerts())

@app.route('/cars')
@login_required
def cars():
    conn = get_connection(); cursor = conn.cursor()
    execute_query(cursor, "SELECT * FROM cars ORDER BY id"); res = cursor.fetchall(); conn.close()
    return render_template('cars.html', cars=res)

@app.route('/drivers')
@login_required
def drivers():
    conn = get_connection(); cursor = conn.cursor()
    execute_query(cursor, "SELECT * FROM drivers ORDER BY id"); res = cursor.fetchall(); conn.close()
    return render_template('drivers.html', drivers=res)

@app.route('/fuel')
@login_required
def fuel():
    conn = get_connection(); cursor = conn.cursor()
    execute_query(cursor, "SELECT f.*, c.name as car_name, d.name as driver_name FROM fuel f JOIN cars c ON f.car_id = c.id LEFT JOIN drivers d ON f.driver_id = d.id ORDER BY f.id DESC LIMIT 50")
    res = cursor.fetchall(); conn.close()
    return render_template('fuel.html', fuel_records=res)

@app.route('/maintenance')
@login_required
def maintenance():
    conn = get_connection(); cursor = conn.cursor()
    execute_query(cursor, "SELECT m.*, c.name as car_name FROM maintenance m JOIN cars c ON m.car_id = c.id ORDER BY m.id DESC LIMIT 100"); res = cursor.fetchall()
    conn.close()
    return render_template('maintenance.html', records=res)

@app.route('/admin/pending')
@login_required
def admin_pending_trips():
    conn = get_connection(); cursor = conn.cursor()
    execute_query(cursor, "SELECT t.*, c.name as car_name, d.name as driver_name FROM trips t JOIN cars c ON t.car_id = c.id JOIN drivers d ON t.driver_id = d.id WHERE t.trip_source = 'driver_app' AND (t.admin_seen IS NULL OR t.admin_seen = 0) ORDER BY t.id DESC")
    res = cursor.fetchall(); conn.close()
    return render_template('admin_pending.html', pending_trips=res)

@app.route('/reports')
@login_required
def reports():
    month = int(request.args.get('month', datetime.now().month)); year = int(request.args.get('year', datetime.now().year))
    conn = get_connection(); cursor = conn.cursor()
    execute_query(cursor, "SELECT * FROM cars ORDER BY name"); cars_list = cursor.fetchall()
    execute_query(cursor, "SELECT * FROM drivers ORDER BY name"); drivers_list = cursor.fetchall()
    execute_query(cursor, "SELECT * FROM requesting_parties ORDER BY name"); parties_list = cursor.fetchall()
    years = list(range(2024, datetime.now().year + 2)); report_data = None
    if request.args.get('month'):
        month_str = f"{year:04d}-{month:02d}"
        execute_query(cursor, "SELECT t.*, c.name as car_name, d.name as driver_name FROM trips t JOIN cars c ON t.car_id = c.id JOIN drivers d ON t.driver_id = d.id WHERE t.trip_date LIKE ?", (month_str + '%',)); trips = cursor.fetchall()
        execute_query(cursor, "SELECT f.*, c.name as car_name FROM fuel f JOIN cars c ON f.car_id = c.id WHERE f.fuel_date LIKE ?", (month_str + '%',)); fuels = cursor.fetchall()
        execute_query(cursor, "SELECT m.*, c.name as car_name FROM maintenance m JOIN cars c ON m.car_id = c.id WHERE m.date LIKE ?", (month_str + '%',)); maints = cursor.fetchall()
        report_data = {'summary': {'total_trips': len(trips), 'total_km': sum(t['distance'] or 0 for t in trips), 'total_fuel_cost': sum(get_effective_cost(f) for f in fuels), 'total_maint_cost_syp': sum(m['cost'] or 0 for m in maints if m.get('currency', 'ل.س') == 'ل.س'), 'total_maint_cost_usd': sum(m['cost'] or 0 for m in maints if m.get('currency') == 'دولار')}}
    conn.close()
    return render_template('reports.html', cars=cars_list, drivers=drivers_list, parties=parties_list, years=years, selected_month=month, selected_year=year, report_data=report_data)

@app.route('/reports/trips')
@login_required
def reports_trips():
    conn = get_connection(); cursor = conn.cursor()
    execute_query(cursor, "SELECT * FROM cars ORDER BY name"); cars_list = cursor.fetchall()
    execute_query(cursor, "SELECT * FROM drivers ORDER BY name"); drivers_list = cursor.fetchall()
    execute_query(cursor, "SELECT * FROM requesting_parties ORDER BY name"); parties_list = cursor.fetchall()
    query = "SELECT t.*, c.name as car_name, d.name as driver_name FROM trips t JOIN cars c ON t.car_id = c.id JOIN drivers d ON t.driver_id = d.id WHERE 1=1"
    params = []
    for k in ['car_id', 'driver_id', 'date_from', 'date_to']:
        v = request.args.get(k)
        if v:
            if k == 'car_id': query += " AND t.car_id = ?"; params.append(v)
            if k == 'driver_id': query += " AND t.driver_id = ?"; params.append(v)
            if k == 'date_from': query += " AND t.trip_date >= ?"; params.append(v)
            if k == 'date_to': query += " AND t.trip_date <= ?"; params.append(v)
    query += " ORDER BY t.id DESC"
    execute_query(cursor, query, params); trips = cursor.fetchall(); conn.close()
    summary = {'total_trips': len(trips), 'total_km': sum(t['distance'] or 0 for t in trips)}
    return render_template('reports_trips.html', cars=cars_list, drivers=drivers_list, parties=parties_list, trips=trips, summary=summary)

@app.route('/reports/fuel')
@login_required
def reports_fuel():
    conn = get_connection(); cursor = conn.cursor()
    query = "SELECT f.*, c.name as car_name, d.name as driver_name FROM fuel f JOIN cars c ON f.car_id = c.id LEFT JOIN drivers d ON f.driver_id = d.id WHERE 1=1"
    execute_query(cursor, query); fuel_records = cursor.fetchall(); conn.close()
    summary = {'total_records': len(fuel_records), 'total_liters': sum(f['liters'] or 0 for f in fuel_records), 'total_cost': sum(get_effective_cost(f) for f in fuel_records)}
    return render_template('reports_fuel.html', fuel_records=fuel_records, summary=summary)

@app.route('/reports/export')
@login_required
def export_excel():
    month = int(request.args.get('month', datetime.now().month)); year = int(request.args.get('year', datetime.now().year))
    month_str = f"{year:04d}-{month:02d}"
    conn = get_connection(); cursor = conn.cursor()
    execute_query(cursor, "SELECT t.*, c.name as car_name, d.name as driver_name FROM trips t JOIN cars c ON t.car_id = c.id JOIN drivers d ON t.driver_id = d.id WHERE t.trip_date LIKE ?", (month_str + '%',)); trips = cursor.fetchall()
    execute_query(cursor, "SELECT * FROM cars ORDER BY name"); cars_list = cursor.fetchall()
    conn.close()
    wb = Workbook(); ws_sum = wb.active; ws_sum.title = "الملخص"; ws_sum.sheet_view.rightToLeft = True
    ws_sum.append(['التقرير الشهري', month_str]); ws_sum.append(['إجمالي الرحلات', len(trips)])
    for c in cars_list:
        c_trips = [t for t in trips if t['car_id'] == c['id']]
        if not c_trips: continue
        ws_c = wb.create_sheet(title=f"سيارة - {c['name'][:20]}"); ws_c.sheet_view.rightToLeft = True
        ws_c.append([f"السيارة: {c['name']}"]); ws_c.append(['التاريخ', 'السائق', 'المسافة'])
        for t in c_trips: ws_c.append([t['trip_date'], t['driver_name'], t['distance'] or 0])
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name=f'qamh_report_{month_str}.xlsx')

# ============== البوت ==============
@app.route('/telegram_webhook', methods=['POST'])
def telegram_webhook():
    data = request.get_json()
    if 'message' in data and data['message'].get('photo'):
        chat_id = data['message']['chat']['id']
        file_id = data['message']['photo'][-1]['file_id']
        try:
            file_info = requests.get(f"https://api.telegram.org/bot{TELEGRAM_TOKEN}/getFile?file_id={file_id}").json()
            img_data = requests.get(f"https://api.telegram.org/file/bot{TELEGRAM_TOKEN}/{file_info['result']['file_path']}").content
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S'); filename = f"tg_{chat_id}_{timestamp}.jpg"
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            with open(filepath, 'wb') as f: f.write(img_data)
            conn = get_connection(); cursor = conn.cursor()
            now = datetime.now() + timedelta(hours=3)
            execute_query(cursor, "INSERT INTO pending_fuel (driver_id, car_id, fuel_date, fuel_time, status, created_at, car_image) VALUES (NULL, NULL, ?, ?, 'pending', ?, ?)", (now.strftime('%Y-%m-%d'), now.strftime('%H:%M:%S'), now.strftime('%Y-%m-%d %H:%M:%S'), filename))
            conn.commit(); conn.close()
            requests.post(f"https://api.telegram.org/bot{TELEGRAM_TOKEN}/sendMessage", json={'chat_id': chat_id, 'text': '✅ تم استلام الصورة'})
        except Exception as e: print(e)
    return jsonify({'status': 'ok'})

@app.route('/admin/pending_fuel')
@login_required
def admin_pending_fuel():
    conn = get_connection(); cursor = conn.cursor()
    execute_query(cursor, "SELECT * FROM pending_fuel WHERE status = 'pending' ORDER BY id DESC")
    pending_fuels = cursor.fetchall(); conn.close()
    return render_template('admin_pending_fuel.html', pending_fuels=pending_fuels)

@app.route('/driver')
def driver_app(): return render_template('driver_app.html')

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)
